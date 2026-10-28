import openpyxl, json, os

files = {
    'matematik': ('meb_data/matematik.xlsx', [('5', '5.sınıf'), ('6', '6.sınıf'), ('7', '7.sınıf'), ('8', '8.sınıf')]),
    'fen_bilimleri': ('meb_data/fen_bilimleri.xlsx', [('5', '5. sınıf'), ('6', '6. sınıf'), ('7', '7. sınıf'), ('8', '8. sınıf')]),
    'turkce': ('meb_data/turkce.xlsx', [('5', '5. Sınıf'), ('6', '6. SINIF'), ('7', '7. SINIF'), ('8', '8. SINIF')]),
    'sosyal_bilgiler': ('meb_data/sosyal_bilgiler.xlsx', [('5', '5. sınıf '), ('6', '6. sınıf '), ('7', '7. sınıf '), ('8', '8. sınıf T.C. inkılap tarihi ')]),
    'din_kulturu': ('meb_data/din_kulturu.xlsx', [('5', '5. Sınıf'), ('6', '6. Sınıf'), ('7', '7. Sınıf'), ('8', '8. Sınıf')]),
    'ingilizce': ('meb_data/ingilizce.xlsx', [('5', '5. sınıf'), ('6', '6. sınıf'), ('7', '7. sınıf'), ('8', '8. sınıf')]),
}

all_curriculum = {}

for subj, (path, grade_sheets) in files.items():
    wb = openpyxl.load_workbook(path, data_only=True)
    all_curriculum[subj] = {}
    
    for grade_str, sh_name in grade_sheets:
        sheet = wb[sh_name]
        grade = int(grade_str)
        
        # Step 1: Find the row that contains "1. Dönem 1."
        header_row_term = None
        for r in range(1, 15):
            for c in range(1, 20):
                val = sheet.cell(r, c).value
                if val and '1. Dönem 1.' in str(val):
                    header_row_term = r
                    break
            if header_row_term:
                break
                
        if not header_row_term:
            print(f'Warning: Could not find 1. Dönem 1. header in {subj} {sh_name}')
            continue
            
        scenario_row = header_row_term + 1
        
        # Step 2: Find which columns belong to 1. Dönem 1. Ortak Yazılı Sınav
        # Find start col of 1. Dönem 1. and start col of 1. Dönem 2.
        term1_start_col = None
        term2_start_col = None
        
        for c in range(1, 30):
            val = sheet.cell(header_row_term, c).value
            if val and '1. Dönem 1.' in str(val) and term1_start_col is None:
                term1_start_col = c
            elif val and '1. Dönem 2.' in str(val) and term2_start_col is None:
                term2_start_col = c
                
        # If term2_start_col was not in header_row_term (sometimes merged or in next row), look for it
        if term2_start_col is None:
            # Look in scenario_row or subsequent cols
            for c in range(term1_start_col + 1, 30):
                val = sheet.cell(header_row_term, c).value
                val_sc = sheet.cell(scenario_row, c).value
                if val and '1. Dönem 2.' in str(val):
                    term2_start_col = c
                    break
                elif val_sc and '1. Dönem 2.' in str(val_sc):
                    term2_start_col = c
                    break
                    
        # Collect scenario names for 1. Dönem 1.
        scenarios = []
        c = term1_start_col
        while c < 30:
            if term2_start_col and c >= term2_start_col:
                break
            sc_val = sheet.cell(scenario_row, c).value
            term_val = sheet.cell(header_row_term, c).value
            if term_val and '1. Dönem 2.' in str(term_val):
                break
            if sc_val and 'Senaryo' in str(sc_val):
                scenarios.append((c, str(sc_val).strip()))
            elif sc_val and '1. Dönem 2.' in str(sc_val):
                break
            elif not sc_val and len(scenarios) > 0:
                # check if next col has Senaryo
                next_val = sheet.cell(scenario_row, c + 1).value
                if not next_val or 'Senaryo' not in str(next_val):
                    break
            c += 1
            
        # Step 3: Parse outcomes and question counts
        outcomes_list = []
        for r in range(scenario_row + 1, sheet.max_row + 1):
            # outcome text is usually in col 2 or 3
            # let's find the text that looks like an outcome
            row_texts = [(c_idx, str(sheet.cell(r, c_idx).value).strip()) for c_idx in range(1, term1_start_col) if sheet.cell(r, c_idx).value is not None]
            
            # identify outcome text
            outcome_text = None
            unit_text = ''
            topic_text = ''
            
            for c_idx, txt in row_texts:
                if any(txt.startswith(prefix) for prefix in ['T.', 'MAT.', 'FB.', 'F.', 'M.', 'SB', 'İNK.', 'İTA.', 'DKAB.', 'E', 'ENG.']) or ('bilme' in txt) or ('anlar' in txt) or ('açıklar' in txt) or ('kavrar' in txt) or ('analiz eder' in txt) or ('çözümler' in txt):
                    outcome_text = txt
                    break
                    
            if not outcome_text and len(row_texts) >= 1:
                # maybe last text in row_texts before numbers
                outcome_text = row_texts[-1][1]
                
            if outcome_text:
                # find unit and topic from earlier cols
                for c_idx, txt in row_texts:
                    if txt != outcome_text:
                        if not unit_text:
                            unit_text = txt
                        elif not topic_text:
                            topic_text = txt
                            
            # get counts for each scenario
            counts = []
            for sc_col, sc_name in scenarios:
                cell_v = sheet.cell(r, sc_col).value
                cnt = int(cell_v) if isinstance(cell_v, (int, float)) else 0
                counts.append(cnt)
                
            if any(counts) and outcome_text:
                outcomes_list.append({
                    'unit': unit_text,
                    'topic': topic_text,
                    'outcome': outcome_text,
                    'counts': counts
                })
                
        all_curriculum[subj][grade] = {
            'scenarios': [sc[1] for sc in scenarios],
            'scenario_cols': [sc[0] for sc in scenarios],
            'outcomes': outcomes_list
        }
        
        # print summary
        sc_names = [sc[1] for sc in scenarios]
        q_totals = [sum(o['counts'][idx] for o in outcomes_list) for idx in range(len(scenarios))]
        print(f'{subj} Grade {grade}: {len(scenarios)} scenarios: {list(zip(sc_names, q_totals))}')

with open('meb_data/middle_school_curriculum.json', 'w', encoding='utf-8') as f:
    json.dump(all_curriculum, f, ensure_ascii=False, indent=2)

print('Saved middle_school_curriculum.json!')
