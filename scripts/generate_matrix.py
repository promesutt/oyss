import openpyxl, json

files = {
    'matematik': ('meb_data/matematik.xlsx', [('5', '5.sınıf'), ('6', '6.sınıf'), ('7', '7.sınıf'), ('8', '8.sınıf')]),
    'fen_bilimleri': ('meb_data/fen_bilimleri.xlsx', [('5', '5. sınıf'), ('6', '6. sınıf'), ('7', '7. sınıf'), ('8', '8. sınıf')]),
    'turkce': ('meb_data/turkce.xlsx', [('5', '5. Sınıf'), ('6', '6. SINIF'), ('7', '7. SINIF'), ('8', '8. SINIF')]),
    'sosyal_bilgiler': ('meb_data/sosyal_bilgiler.xlsx', [('5', '5. sınıf '), ('6', '6. sınıf '), ('7', '7. sınıf '), ('8', '8. sınıf T.C. inkılap tarihi ')]),
    'din_kulturu': ('meb_data/din_kulturu.xlsx', [('5', '5. Sınıf'), ('6', '6. Sınıf'), ('7', '7. Sınıf'), ('8', '8. Sınıf')]),
    'ingilizce': ('meb_data/ingilizce.xlsx', [('5', '5. sınıf'), ('6', '6. sınıf'), ('7', '7. sınıf'), ('8', '8. sınıf')]),
}

res = {}

for subj, (path, grade_sheets) in files.items():
    wb = openpyxl.load_workbook(path, data_only=True)
    res[subj] = {}
    
    for grade_str, sh_name in grade_sheets:
        grade = int(grade_str)
        sheet = wb[sh_name]
        
        # Special handling for 6.sınıf matematik (Ülke Geneli Sınav)
        if subj == 'matematik' and grade == 6:
            res[subj][grade] = {
                'scenarios': ['Senaryo 1 (Ülke Geneli Formatı - 6 Soru)', 'Senaryo 2 (Beceri ve Analiz - 5 Soru)', 'Senaryo 3 (Kapsamlı Senaryo - 7 Soru)'],
                'outcomes': [
                    {
                        'unit': 'SAYILAR VE NİCELİKLER (1)',
                        'topic': 'Bir Doğal Sayının Çarpanları ve Katları',
                        'code': 'MAT.6.1.1',
                        'title': 'Karşılaştığı problem durumlarında bir doğal sayının çarpanlarını ve katlarını belirleyebilme',
                        'counts': [1, 1, 2]
                    },
                    {
                        'unit': 'SAYILAR VE NİCELİKLER (1)',
                        'topic': 'Bölünebilme Kriterleri',
                        'code': 'MAT.6.1.2',
                        'title': 'Bir doğal sayının 2, 3, 4, 5, 6, 9 ve 10\'a kalansız bölünebilme kurallarını açıklayabilme ve kullanabilme',
                        'counts': [2, 1, 2]
                    },
                    {
                        'unit': 'SAYILAR VE NİCELİKLER (1)',
                        'topic': 'Asal Sayılar ve Asal Çarpanlar',
                        'code': 'MAT.6.1.3',
                        'title': 'Bir doğal sayının asal olma durumunu ve asal çarpanlarını belirleyebilme',
                        'counts': [1, 1, 1]
                    },
                    {
                        'unit': 'SAYILAR VE NİCELİKLER (1)',
                        'topic': 'Bir Doğal Sayının Çarpanları ve Katları',
                        'code': 'MAT.6.1.4',
                        'title': 'Günlük hayat problemleri ya da matematiksel durumlarda iki doğal sayının ortak bölenleri ve katlarını yorumlayabilme',
                        'counts': [2, 2, 2]
                    }
                ]
            }
            continue

        # Find header row
        h_row = None
        for r in range(1, 15):
            for c in range(1, 20):
                v = sheet.cell(r, c).value
                if v and '1. Dönem 1.' in str(v):
                    h_row = r
                    break
            if h_row:
                break
                
        sc_row = h_row + 1
        
        # Find 1. Donem 1. start and 1. Donem 2. start
        t1_col = None
        t2_col = None
        for c in range(1, 30):
            v5 = sheet.cell(h_row, c).value
            v6 = sheet.cell(sc_row, c).value
            if v5 and '1. Dönem 1.' in str(v5) and t1_col is None:
                t1_col = c
            if (v5 and '1. Dönem 2.' in str(v5)) or (v6 and '1. Dönem 2.' in str(v6)):
                if t2_col is None and c > (t1_col or 0):
                    t2_col = c
                    
        scenarios = []
        c = t1_col
        while c < (t2_col or 25):
            sc_val = sheet.cell(sc_row, c).value
            term_val = sheet.cell(h_row, c).value
            if term_val and '1. Dönem 2.' in str(term_val):
                break
            if sc_val and 'Senaryo' in str(sc_val):
                scenarios.append((c, str(sc_val).strip()))
            elif sc_val and '1. Dönem 2.' in str(sc_val):
                break
            c += 1
            
        outcomes = []
        for r in range(sc_row + 1, sheet.max_row + 1):
            counts = [int(sheet.cell(r, sc_col).value) if isinstance(sheet.cell(r, sc_col).value, (int, float)) else 0 for sc_col, _ in scenarios]
            if not any(counts):
                continue
                
            # find outcome text and code
            texts = [str(sheet.cell(r, c_idx).value).strip() for c_idx in range(1, t1_col) if sheet.cell(r, c_idx).value is not None]
            
            raw_outcome = texts[-1] if texts else ''
            # check if there's an explicit outcome code
            code = ''
            title = raw_outcome
            for t in texts:
                # check for pattern like MAT.5.1.1 or F.5.1.1.1 or SB.5.1.1 or T.5.1.1 etc.
                parts = t.split()
                first = parts[0].rstrip('.')
                if any(first.startswith(p) for p in ['T.', 'MAT.', 'FB.', 'F.', 'M.', 'SB', 'İNK', 'İTA', 'DKAB', 'E', 'ENG']):
                    code = first
                    title = ' '.join(parts[1:])
                    break
                    
            if not code:
                # generate standard code
                prefix = {'matematik': 'MAT', 'fen_bilimleri': 'FB', 'turkce': 'T', 'sosyal_bilgiler': 'SB' if grade < 8 else 'İNK', 'din_kulturu': 'DKAB', 'ingilizce': 'E'}[subj]
                code = f'{prefix}.{grade}.{len(outcomes)+1}'
                
            unit = texts[0] if len(texts) > 1 else ''
            topic = texts[1] if len(texts) > 2 else unit
            
            outcomes.append({
                'unit': unit,
                'topic': topic,
                'code': code,
                'title': title,
                'counts': counts
            })
            
        res[subj][grade] = {
            'scenarios': [sc[1] for sc in scenarios],
            'outcomes': outcomes
        }

with open('meb_data/clean_middle_school_matrix.json', 'w', encoding='utf-8') as f:
    json.dump(res, f, ensure_ascii=False, indent=2)

print('clean_middle_school_matrix.json generated successfully!')
for s in res:
    for g in res[s]:
        sc_len = len(res[s][g]['scenarios'])
        o_len = len(res[s][g]['outcomes'])
        q_sum = [sum(o['counts'][i] for o in res[s][g]['outcomes']) for i in range(sc_len)]
        print(f'{s} {g}. Sınıf: {sc_len} scenarios ({q_sum}), {o_len} outcomes')
